# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

out=Path('千抽福利活动需求文档.docx')
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.7); sec.bottom_margin=Inches(0.7); sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.8)
styles=doc.styles
styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5)
for name,size in [('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
    st=styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def borders(table):
    tblPr=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)
    tblPr.append(b)
def p(text='',style=None,bold=False):
    para=doc.add_paragraph(style=style); run=para.add_run(text); run.bold=bold; return para

t=doc.add_paragraph(style='Title'); t.alignment=WD_ALIGN_PARAGRAPH.CENTER; t.add_run('千抽福利活动需求文档')
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=sub.add_run('章节通关领取活动'); r.font.size=Pt(11); r.font.color.rgb=RGBColor(89,89,89)
p('本文档定义千抽福利活动的章节领取规则、页面展示、列表定位、连续领取和验收要求。')

doc.add_heading('1 文档属性', level=1)
tbl=doc.add_table(rows=1, cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(tbl)
for i,h in enumerate(['属性','内容']): shade(tbl.rows[0].cells[i],'1F4E78'); rr=tbl.rows[0].cells[i].paragraphs[0].add_run(h); rr.bold=True; rr.font.color.rgb=RGBColor(255,255,255)
for a,b in [('模块','千抽福利活动'),('需求类型','限时活动页面与奖励领取'),('来源','用户提供的活动界面图及已确认规则'),('版本状态','需求稿')]:
    cells=tbl.add_row().cells; cells[0].text=a; cells[1].text=b

doc.add_heading('2 功能概述与目标', level=1)
p('玩家通过推进主线章节获得奖励领取资格。每个章节对应一个千抽奖励，玩家通关指定章节后即可领取。活动页面以纵向章节奖励列表呈现，帮助玩家查看当前进度、可领取奖励和后续目标。')

doc.add_heading('3 核心规则与状态', level=1)
doc.add_heading('3.1 章节领取条件', level=2)
p('章节奖励的领取条件为通关指定章节。玩家通关对应章节后，该章节奖励进入可领取状态。')
doc.add_heading('3.2 奖励状态', level=2)
for title,desc in [('未解锁','尚未通关对应指定章节，奖励不可领取。'),('可领取','已通关对应指定章节，奖励尚未领取。'),('已领取','奖励已成功发放，不能重复领取。')]:
    para=doc.add_paragraph(); para.add_run(title+'：').bold=True; para.add_run(desc)
doc.add_heading('3.3 连续领取', level=2)
p('当多个章节奖励同时满足领取条件时，玩家点击一次“领取”按钮后，系统连续领取所有已满足且未领取的章节奖励，并刷新各章节的状态。')

doc.add_heading('4 界面需求', level=1)
doc.add_heading('4.1 活动页面', level=2)
doc.add_heading('4.1.1 显示内容', level=3)
for s in ['顶部区域：活动主题与规则说明。','章节奖励区域：按章节顺序纵向展示章节名称、奖励内容和领取状态。','按钮区域：显示“领取”按钮。']:
    doc.add_paragraph(s, style='List Bullet')
doc.add_heading('4.1.2 显示规则', level=3)
p('章节奖励区域支持上下滑动。奖励按章节顺序从上到下排列。未解锁、可领取和已领取奖励使用不同状态表现，具体视觉样式沿用活动界面设计。')
doc.add_heading('4.1.3 列表定位', level=3)
for s in ['玩家尚未领取任何奖励时：进入页面后，章节奖励区域保持初始位置。','玩家已有领取记录时：进入页面后，将历史领取记录中章节序号最大的奖励定位在可视区域中间。','连续领取完成后：刷新奖励状态，并保持最远已领取奖励位于可视区域中间。']:
    doc.add_paragraph(s, style='List Bullet')
doc.add_heading('4.1.4 按钮规则', level=3)
for s in ['存在至少一个已满足且未领取的章节奖励时，“领取”按钮可点击。','不存在可领取奖励时，“领取”按钮不可用。','点击“领取”后，连续领取所有已满足且未领取的章节奖励。']:
    doc.add_paragraph(s, style='List Bullet')

doc.add_heading('5 交互流程', level=1)
flow=[('进入活动页','读取玩家章节进度和奖励领取记录。'),('确定初始位置','无领取记录则使用初始位置；有领取记录则将最远已领取奖励定位到区域中间。'),('查看奖励','玩家上下滑动奖励区域，查看各章节状态。'),('领取奖励','玩家点击“领取”，系统连续发放所有已满足且未领取的章节奖励。'),('刷新页面','更新奖励状态、领取按钮状态和最远已领取章节定位。')]
tbl=doc.add_table(rows=1, cols=2); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(tbl)
for i,h in enumerate(['步骤','行为']): shade(tbl.rows[0].cells[i],'1F4E78'); rr=tbl.rows[0].cells[i].paragraphs[0].add_run(h); rr.bold=True; rr.font.color.rgb=RGBColor(255,255,255)
for a,b in flow:
    cells=tbl.add_row().cells; cells[0].text=a; cells[1].text=b

doc.add_heading('6 配置与数据规则', level=1)
p('本次需求确认的配置关系如下：')
nt=doc.add_table(rows=1, cols=3); nt.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(nt)
for i,h in enumerate(['配置对象','必需信息','使用规则']): shade(nt.rows[0].cells[i],'1F4E78'); rr=nt.rows[0].cells[i].paragraphs[0].add_run(h); rr.bold=True; rr.font.color.rgb=RGBColor(255,255,255)
for row in [('章节奖励','章节编号、指定通关章节、奖励内容','用于生成章节奖励列表并判断领取资格'),('玩家进度','主线章节通关进度','判断章节奖励是否满足领取条件'),('领取记录','已领取章节记录','判断奖励状态并确定进入页面时的列表定位')]:
    cells=nt.add_row().cells
    for i,v in enumerate(row): cells[i].text=v

doc.add_heading('7 依赖与职责', level=1)
for s in ['客户端：展示活动页面、章节奖励状态、上下滑动区域、列表定位和领取按钮状态。','服务端：提供玩家章节进度，校验章节领取条件，发放奖励并保存领取记录。','配置：提供章节奖励顺序、指定通关章节和奖励内容。']:
    doc.add_paragraph(s, style='List Bullet')

doc.add_heading('8 验收要求', level=1)
checks=['未领取任何奖励时，进入页面后奖励列表处于初始位置。','已领取至少一个奖励时，进入页面后最远已领取奖励位于可视区域中间。','玩家通关指定章节后，对应奖励显示为可领取状态。','同时满足多个章节条件时，点击一次领取按钮可连续领取多个奖励。','已领取奖励不会再次发放。','连续领取完成后，各章节状态和领取按钮状态正确刷新。','没有可领取奖励时，领取按钮不可用。','奖励列表可正常上下滑动，滑动不会改变章节奖励状态。']
for i,s in enumerate(checks,1): doc.add_paragraph(f'{i}. {s}')

# footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('千抽福利活动需求文档').font.size=Pt(9)

doc.save(out)
print(out.resolve())

