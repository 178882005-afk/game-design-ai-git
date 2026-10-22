"""Read-only structural and content validation of the merged activity workbook."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree as ET
from openpyxl import load_workbook

from merge_activity_tabs import SOURCE, OUTPUT, SHEET_NAME, M, R, P, CT, Q


def cells(root):
    return {cell.get("r"): cell for cell in root.iter(Q(M, "c"))}


def main() -> None:
    with ZipFile(SOURCE) as before, ZipFile(OUTPUT) as after:
        assert after.testzip() is None
        old_parts = set(before.namelist())
        new_parts = set(after.namelist())
        assert len(old_parts) == 174 and len(new_parts) == 176, (len(old_parts), len(new_parts))
        removed = old_parts - new_parts
        added = new_parts - old_parts
        assert removed == {f"xl/worksheets/sheet{i}.xml" for i in range(36, 42)}
        assert added == {f"xl/tables/table{i}.xml" for i in range(53, 60)} | {"xl/worksheets/_rels/sheet35.xml.rels"}
        permitted_changes = {"xl/worksheets/sheet35.xml", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "docProps/app.xml"}
        changed = {name for name in old_parts & new_parts if before.read(name) != after.read(name)}
        assert changed == permitted_changes, changed
        for i in range(1, 35):
            name = f"xl/worksheets/sheet{i}.xml"
            assert before.read(name) == after.read(name)
        for name in old_parts & new_parts - permitted_changes:
            assert before.read(name) == after.read(name), name

        book = ET.fromstring(after.read("xl/workbook.xml"))
        sheet_nodes = book.find(Q(M, "sheets"))
        names = [node.get("name") for node in sheet_nodes]
        assert len(names) == 35 and names[-1] == SHEET_NAME
        assert all(not name.startswith("现行") for name in names)
        assert len(names) == len(set(names))

        merged = ET.fromstring(after.read("xl/worksheets/sheet35.xml"))
        merged_cells = cells(merged)
        assert len(merged.findall(f"{Q(M, 'sheetData')}/{Q(M, 'row')}")) == 133
        assert merged.find(Q(M, "dimension")).get("ref") == "A1:Z145"
        table_refs = ["A4:J7", "A13:Z22", "A28:N48", "A54:J61", "A67:L69", "A75:J84", "A90:J145"]
        original_lengths = [7, 13, 24, 11, 6, 13, 59]
        offsets = [0, 9, 24, 50, 63, 71, 86]
        formula_total = 0
        for i, offset in zip(range(35, 42), offsets):
            original = ET.fromstring(before.read(f"xl/worksheets/sheet{i}.xml"))
            source_cells = cells(original)
            for ref, cell in source_cells.items():
                new_ref = re.sub(r"\d+$", lambda m: str(int(m.group()) + offset), ref)
                new_cell = merged_cells[new_ref]
                before_formula = cell.find(Q(M, "f"))
                after_formula = new_cell.find(Q(M, "f"))
                if before_formula is not None:
                    assert after_formula is not None, (i, ref)
                    formula_total += 1
                else:
                    assert after_formula is None, (i, ref)
                # Source style, type, and cached literal values remain unchanged.
                assert cell.get("s") == new_cell.get("s"), (i, ref)
                assert cell.get("t") == new_cell.get("t"), (i, ref)
                source_value = cell.find(Q(M, "v"))
                new_value = new_cell.find(Q(M, "v"))
                assert (source_value.text if source_value is not None else None) == (new_value.text if new_value is not None else None), (i, ref)
                source_inline = cell.find(Q(M, "is"))
                new_inline = new_cell.find(Q(M, "is"))
                assert ("".join(source_inline.itertext()) if source_inline is not None else None) == ("".join(new_inline.itertext()) if new_inline is not None else None), (i, ref)
        assert formula_total == 128
        for row in range(76, 85):
            formula = merged_cells[f"F{row}"].find(Q(M, "f")).text
            assert f"E{row}" in formula and "$E$76:$E$84" in formula
            assert "$E$5:$E$13" not in formula

        tableparts = merged.find(Q(M, "tableParts"))
        assert tableparts is not None and len(tableparts) == 7
        sheetrels = ET.fromstring(after.read("xl/worksheets/_rels/sheet35.xml.rels"))
        assert len(sheetrels) == 7
        for table_no, table_ref in zip(range(53, 60), table_refs):
            table = ET.fromstring(after.read(f"xl/tables/table{table_no}.xml"))
            assert table.get("id") == str(table_no + 23), table_no
            assert table.get("ref") == table_ref
            assert table.find(Q(M, "autoFilter")).get("ref") == table_ref
            cols = table.find(Q(M, "tableColumns"))
            assert len(cols) == int(cols.get("count"))
            header_row = int(table_ref.split(":")[0][1:])
            for column_index, column in enumerate(cols, start=1):
                letters = ""
                n = column_index
                while n:
                    n, rem = divmod(n - 1, 26)
                    letters = chr(65 + rem) + letters
                cell = merged_cells[f"{letters}{header_row}"]
                assert "".join(cell.find(Q(M, "is")).itertext()) == column.get("name")

    # A second, independent parser checks that Excel will recognize the tables.
    wb = load_workbook(OUTPUT, read_only=False, data_only=False, keep_links=True)
    assert len(wb.sheetnames) == 35
    ws = wb[SHEET_NAME]
    assert len(ws.tables) == 7
    assert Counter(t.ref for t in ws.tables.values()) == Counter(table_refs)
    assert ws.max_row == 145 and ws.max_column == 26
    assert ws["F76"].value == "=ROUND(E76/SUM($E$76:$E$84)*100,2)"
    assert ws["N14"].value.startswith("=IF(L14=")
    assert ws["I91"].value.startswith("=IF(E91=")
    wb.close()
    print("PASS: ZIP, 35 sheets, 7 filter tables, 105 data records, 128 formulas, all old parts intact")
    print("OUTPUT", OUTPUT)


if __name__ == "__main__":
    main()
