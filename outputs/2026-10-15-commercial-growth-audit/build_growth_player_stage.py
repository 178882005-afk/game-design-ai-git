"""Read-only extraction of the player, stamina and mainline growth audit."""

from __future__ import annotations

import csv
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl


ROOT = Path(r"E:\x1\配置\trunk")
NUM = ROOT / "策划文档" / "数值文档"
XML = ROOT / "xml_gy" / "normal"
ACTUAL = Path(
    r"D:\Administrator\Desktop\线上数据复盘\分析任务"
    r"\20261008_小游戏渠道主线通关分析\reports"
    r"\梦境护卫队_小游戏主线复盘_20261009_V3.xlsx"
)
OLD_RESOURCE = Path(r"D:\Administrator\Desktop\梦境护卫队资源产销7-8月.xlsx")
P90_RESOURCE = Path(r"D:\Administrator\Desktop\梦境护卫队_每日P90单玩家钻石来源_20261015.xlsx")
OUT = Path(__file__).with_name("growth_player_stage.tsv")
NS = {"s": "urn:schemas-microsoft-com:office:spreadsheet"}
INDEX = "{urn:schemas-microsoft-com:office:spreadsheet}Index"

FIELDS = [
    "主题", "日龄D", "人群或目标", "指标", "当前规划", "较旧规划",
    "正式配置", "线上实测", "单位", "规划证据", "配置证据", "实测证据", "判读",
    "press_days配置值", "主线索引开战是否按此限制",
]


def xml_rows(filename: str, sheet_name: str):
    tree = ET.parse(XML / filename)
    for ws in tree.findall(".//s:Worksheet", NS):
        if ws.attrib.get("{urn:schemas-microsoft-com:office:spreadsheet}Name") != sheet_name:
            continue
        row_num = 0
        for row in ws.findall("./s:Table/s:Row", NS):
            row_num = int(row.attrib.get(INDEX, row_num + 1))
            col_num = 0
            values = {}
            for cell in row.findall("s:Cell", NS):
                col_num = int(cell.attrib.get(INDEX, col_num + 1))
                data = cell.find("s:Data", NS)
                if data is not None and data.text is not None:
                    values[col_num] = data.text
            yield row_num, values
        break


def add(rows, topic, day, group, metric, current="", old="", config="", actual="",
        unit="", plan_source="", config_source="", actual_source="", note="",
        gate="", relationship=""):
    rows.append([topic, day, group, metric, current, old, config, actual, unit,
                 plan_source, config_source, actual_source, note, gate, relationship])


def main():
    rows = []
    old = openpyxl.load_workbook(NUM / "《梦境护卫队》1-0数值总表.xlsx", read_only=True, data_only=True)
    old_stage = {int(r[3].value): int(r[4].value)
                 for r in old["养成框架"].iter_rows(min_row=64, max_row=77, min_col=1, max_col=5)
                 if isinstance(r[3].value, (int, float)) and isinstance(r[4].value, (int, float))}

    growth = openpyxl.load_workbook(NUM / "《梦境护卫队》1-1玩家等级.xlsx", read_only=True, data_only=True)
    expectation = growth["玩家预期"]
    growth_days = {}
    for row_num, r in enumerate(expectation.iter_rows(min_row=3, max_row=35, min_col=1, max_col=14), 3):
        if isinstance(r[2].value, (int, float)):
            growth_days[int(r[2].value)] = (row_num, r)

    design = openpyxl.load_workbook(NUM / "《梦境护卫队》2-0养成预期.xlsx", read_only=True, data_only=True)
    overview = design["总览"]
    supply = openpyxl.load_workbook(NUM / "《梦境护卫队》3-0产出总览.xlsx", read_only=True, data_only=True)
    summary = supply["产销汇总"]
    stage_plan = {int(summary.cell(r, 3).value): r for r in range(13, 32)}

    formal_stage = {}
    for row_num, v in xml_rows("Z_主线推图_chapter_pve_data.xml", "章节配置"):
        if v.get(4) == "主线" and v.get(1, "").isdigit():
            formal_stage[int(v[1])] = (row_num, v)
    formal_level = {}
    for row_num, v in xml_rows("W_玩家等级_leader_data.xml", "等级"):
        if v.get(1, "").isdigit():
            formal_level[int(v[1])] = (row_num, v)

    for day in range(1, 15):
        sr = stage_plan[day]
        for col, group in [(4, "非R"), (5, "小R"), (6, "中R"), (7, "大R")]:
            target = summary.cell(sr, col).value
            old_target = old_stage.get(day, "") if group == "非R" else ""
            frow = formal_stage.get(int(target)) if isinstance(target, int) else None
            gate = frow[1].get(11, "") if frow else ""
            relation = "否；运行服待验"
            gate_note = "；press_days数值晚于目标日，但不能据此判为不可达" if gate and int(gate) > day else ""
            add(rows, "主线", day, group, "计划最高通关", target, old_target, "", "", "关",
                f"3-0产出总览/产销汇总!{openpyxl.utils.get_column_letter(col)}{sr}"
                + (f"；1-0数值总表/养成框架!D{day+63}:E{day+63}" if old_target != "" and day <= 14 else ""),
                (f"Z_主线推图/章节配置!{frow[0]}；DreamGuard Knowledge索引：server/"
                 "src/mod/chapter_pve/role_chapter_pve.erl:L1028-L1108") if frow else "", "",
                "主线索引开战路径未把press_days作为开服日门槛；仍检查功能、战力、前置、体力"
                + gate_note + "。实际服未验证。",
                gate=gate, relationship=relation)

    for day in [1, 2, 3, 7, 14]:
        row_num, r = growth_days[day]
        for exp_col, level_col, group in [(5, 11, "非R"), (6, 12, "小R"), (7, 13, "中R"), (8, 14, "大R")]:
            exp_plan = r[exp_col - 1].value
            level_plan = r[level_col - 1].value
            level_record = formal_level.get(level_plan)
            add(rows, "玩家等级", day, group, "计划等级", level_plan, "", "", "", "级",
                f"1-1玩家等级/玩家预期!{openpyxl.utils.get_column_letter(level_col)}{row_num}",
                "", "", "等级预期是旧策划模型；没有同龄同渠道等级实测。")
            add(rows, "玩家等级", day, group, "计划累计经验与达标阈值", exp_plan, "",
                int(level_record[1][3]) if level_record else "", "", "经验",
                f"1-1玩家等级/玩家预期!{openpyxl.utils.get_column_letter(exp_col)}{row_num}",
                f"W_玩家等级/等级!C{level_record[0]}" if level_record else "", "",
                f"正式配置值为升到计划等级{level_plan}的最低累计经验；策划经验与等级需按配置重算，非线上获得值。")

    for row_num, group in [(29, "非R"), (30, "小R"), (31, "中R"), (32, "大R"), (33, "超R")]:
        add(rows, "体力", "日均假设", group, "分层日预算", overview.cell(row_num, 4).value,
            summary.cell(6, {"非R": 4, "小R": 5, "中R": 6, "大R": 7}.get(group)).value if group != "超R" else "",
            "", "", "点/日", f"2-0养成预期/总览!D{row_num}；3-0产出总览/产销汇总!D6:G6",
            "", "", "两张策划表为不同预算口径；大R 560与340冲突，未选定实际体力来源组合。")
        add(rows, "钻石", "日均假设", group, "分层日预算", overview.cell(row_num, 12).value,
            "", "", "", "钻/日", f"2-0养成预期/总览!L{row_num}", "", "",
            "r65300最新规划数；不是同层线上日均免费钻产出，也不等于付费充值。")

    for row_num, name in [(9, "时间恢复"), (10, "广告全看"), (11, "食堂全领"),
                          (12, "免费日礼包"), (13, "付费月卡"), (14, "特惠1"),
                          (15, "特惠2"), (16, "特惠3")]:
        add(rows, "体力", "日均假设", name, "来源预算", growth["玩家升级"].cell(row_num, 7).value,
            "", "", "", "点/日", f"1-1玩家等级/玩家升级!D{row_num}:K{row_num}", "", "",
            "领取、广告或购买是条件项；只取单项G列，勿再加H列累计。")
    add(rows, "体力", "日均假设", "无溢出自然恢复", "理论上限", "", "", 240, "", "点/日",
        "", "W_玩家等级/等级!E3:F302", "",
        "正式配置为360秒恢复1点、自动恢复上限60；240仅为24小时不溢出的算术上限，不是实际到账。")
    add(rows, "体力", "创建时", "新角色", "初始体力策划假设", 60, "", "", "", "点",
        "1-1玩家等级/玩家升级!F8", "", "", "一次性起始库存，不并入每天体力产出。")

    for target in [10, 30, 40, 57, 60, 100]:
        frow, v = formal_stage[target]
        add(rows, "主线配置", "", f"至第{target}关", "全首通配置门票合计", "", "",
            sum(int(formal_stage[i][1].get(9, "0")) for i in range(1, target + 1)), "", "点",
            "", f"Z_主线推图/章节配置!I9:I{frow}", "",
            "首通各1次的配置成本求和；索引服务端在开战前扣体力，失败需重试则增加成本；运行服未核。前100关均10点/关。")
        add(rows, "主线配置", "", f"第{target}关", "开服日压制", "", "",
            v.get(11, ""), "", "天", "", f"Z_主线推图/章节配置!K{frow}", "",
            "按配置字段原义记录，不代表该日必可通关。")
        add(rows, "主线配置", "", f"第{target}关", "所需战力字段", "", "",
            v.get(10, ""), "", "战力", "", f"Z_主线推图/章节配置!J{frow}", "",
            "配置字段值；与实际战斗胜率和运行服硬性校验未闭合。")

    gold = supply["金币平衡"]
    for r in [7, 8, 9, 10, 14, 15, 16, 17, 19, 20, 21, 22]:
        name = gold.cell(r, 2).value
        amount = gold.cell(r, 3).value
        freq = gold.cell(r, 4).value
        add(rows, "金币", "模型", name, "规划来源候选", amount if isinstance(amount, (int, float)) else "",
            "", "", "", "模型产出价值/" + ("日" if freq == "每日" else "次"),
            f"3-0产出总览/金币平衡!B{r}:E{r}", "", "",
            f"原表频率={freq}，比例列={gold.cell(r,5).value}；多处#N/A，且入口可达性和领取未核，严禁求和作净产。")
    add(rows, "金币", "模型", "全线", "日净产与完成天数", "", "", "", "", "金币/日",
        "3-0产出总览/金币平衡!B3:E22；图纸产出设计!B3:E24", "", "",
        "计算链有#N/A，金币本入口配置create_day=999；缺同一玩家金币库存和养成支出流水，当前不可闭合。")

    live = openpyxl.load_workbook(ACTUAL, read_only=True, data_only=True)
    dist = live["最高关卡分布"]
    observed = []
    for r in dist.iter_rows(min_row=7, max_row=409, min_col=1, max_col=12):
        if isinstance(r[0].value, (int, float)) and isinstance(r[11].value, (int, float)):
            observed.append((int(r[0].value), int(r[11].value), int(r[1].value or 0)))
    total = sum(n for _, n, _ in observed)
    nonr_total = sum(n for _, _, n in observed)
    assert total == 11283 and nonr_total == 10337
    for threshold in [1, 5, 10, 20, 30, 40, 57, 60]:
        for group, denominator, position in [("全体", total, 1), ("非R", nonr_total, 2)]:
            numerator = sum(r[position] for r in observed if r[0] >= threshold)
            add(rows, "主线实绩", "窗口累计", group, f"最高成功通关≥{threshold}", "", "", "",
                f"{numerator}/{denominator}={numerator / denominator:.2%}", "角色/占比", "", "",
                "小游戏主线复盘V3/最高关卡分布!A7:L409；口径与质量!B8:B24",
                "2026-09-24至10-09 16:19混合建角日存量；不与任何D1/D7/D14规划作同日达标率。")
    ranges = live["主线关卡通关情况"]
    for r in [7, 8, 9, 12, 13]:
        label = ranges.cell(r, 1).value
        add(rows, "主线实绩", "窗口累计", label, "场次胜率", "", "", "",
            ranges.cell(r, 6).value, "胜场/场次", "", "",
            f"小游戏主线复盘V3/主线关卡通关情况!A{r}:F{r}",
            "窗口内所有该关段场次比，不是单次玩家通关概率；关段样本不同。")

    old_resource = openpyxl.load_workbook(OLD_RESOURCE, read_only=True, data_only=True)["结果明细"]
    p90_resource = openpyxl.load_workbook(P90_RESOURCE, read_only=True, data_only=True)["P90玩家样本"]
    for name, value, source in [
        ("钻石产出均值", old_resource["L6485"].value, "7-8月资源产销/结果明细!L6485"),
        ("钻石产出P50", old_resource["O6485"].value, "7-8月资源产销/结果明细!O6485"),
        ("钻石产出P90重跑", p90_resource["I2056"].value, "10-15 P90玩家样本!G2056:I2056"),
        ("钻石消耗均值", old_resource["L6515"].value, "7-8月资源产销/结果明细!L6515"),
        ("钻石消耗P50", old_resource["O6515"].value, "7-8月资源产销/结果明细!O6515"),
    ]:
        add(rows, "钻石历史实绩", 1, "2026-07纯新增/全部R档", name, "", "", "", value,
            "钻/角色", "", "", source,
            "历史7月队列，非当前付费层，同一日P90玩家会随日期变化；不能直接匹配2-0分层规划。")

    with OUT.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter="\t")
        writer.writerow(FIELDS)
        writer.writerows(rows)
    print(f"{OUT}: {len(rows)} data rows")


if __name__ == "__main__":
    main()
