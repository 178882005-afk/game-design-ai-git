"""Read selected planning-workbook cells without editing the SVN source files."""

import argparse
import glob
import os
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

ROOT = r"E:\x1\配置\trunk\策划文档\数值文档"
SS = "{urn:schemas-microsoft-com:office:spreadsheet}"


def rows_from_spreadsheetml(path, sheet):
    root = ET.parse(path).getroot()
    for ws in root.findall(SS + "Worksheet"):
        if ws.get(SS + "Name") != sheet:
            continue
        excel_row = 0
        for row in ws.find(SS + "Table").findall(SS + "Row"):
            excel_row = int(row.get(SS + "Index") or excel_row + 1)
            excel_col = 0
            cells = {}
            for cell in row.findall(SS + "Cell"):
                excel_col = int(cell.get(SS + "Index") or excel_col + 1)
                data = cell.find(SS + "Data")
                if data is not None and data.text not in (None, ""):
                    cells[excel_col] = data.text
            yield excel_row, cells


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file_contains")
    ap.add_argument("sheet", nargs="?")
    ap.add_argument("--first", type=int, default=1)
    ap.add_argument("--last", type=int, default=35)
    ap.add_argument("--cols", type=int, default=20)
    ap.add_argument("--all-sheets", action="store_true")
    ap.add_argument("--root", default=ROOT)
    args = ap.parse_args()
    paths = [p for p in glob.glob(os.path.join(args.root, "*" + args.file_contains + "*")) if not os.path.basename(p).startswith("~$")]
    if len(paths) != 1:
        raise SystemExit(f"Expected one workbook; found {paths}")
    path = paths[0]
    if zipfile.is_zipfile(path):
        wb = load_workbook(path, read_only=True, data_only=True)
        if args.all_sheets:
            for ws in wb:
                print(ws.title, ws.max_row, ws.max_column)
            return
        ws = wb[args.sheet]
        for row in ws.iter_rows(min_row=args.first, max_row=args.last, max_col=args.cols):
            values = [f"{cell.coordinate}:{str(cell.value)[:120]}" for cell in row if cell.value is not None]
            if values:
                print(" | ".join(values))
    else:
        root = ET.parse(path).getroot()
        if args.all_sheets:
            for ws in root.findall(SS + "Worksheet"):
                print(ws.get(SS + "Name"))
            return
        for r, cells in rows_from_spreadsheetml(path, args.sheet):
            if args.first <= r <= args.last:
                values = [f"{get_column_letter(c)}{r}:{str(value)[:120]}" for c, value in cells.items() if c <= args.cols]
                if values:
                    print(" | ".join(values))


if __name__ == "__main__":
    main()
