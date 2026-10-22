"""Print selected numeric planning cells without mutating the SVN working copy."""

from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(r"E:\x1\配置\trunk\策划文档\数值文档")


def main() -> None:
    prefix, sheet_name, begin, end, width = sys.argv[1:]
    path = next(ROOT.glob(f"*{prefix}*"))
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=int(begin), max_row=int(end), max_col=int(width)):
        parts = []
        for cell in row:
            if cell.value is None:
                continue
            value = str(cell.value).replace("\n", " ")
            if len(value) > 105:
                value = value[:102] + "..."
            parts.append(f"{cell.coordinate}={value}")
        if parts:
            print(" | ".join(parts))


if __name__ == "__main__":
    main()
